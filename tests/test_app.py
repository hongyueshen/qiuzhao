import io
import json
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch

import requests
from openpyxl import load_workbook

import app


class TrackerTests(unittest.TestCase):
    def test_structured_job_and_requirements(self):
        html = '<script type="application/ld+json">' + json.dumps({'@type': 'JobPosting', 'title': '品牌设计师', 'hiringOrganization': {'name': '示例公司'}, 'description': '<p>岗位职责：开发 Python 工具</p><p>任职要求：熟悉 Figma 与品牌设计，有视觉系统经验</p><p>福利待遇：沟通培训</p>'}) + '</script>'
        job = app.parse_job(html=html)
        self.assertEqual(job['title'], '品牌设计师')
        self.assertEqual(job['company'], '示例公司')
        self.assertIn('Figma', job['keywords'])
        self.assertNotIn('Python', job['keywords'])
        self.assertNotIn('沟通', job['keywords'])

    def test_aliases_and_evidence(self):
        result = app.match('任职要求：Brand identity、Figma、Blender', '简历：品牌识别与 Figma 经验')
        self.assertEqual(result['score'], 67)
        self.assertEqual(result['missing'], ['Blender'])
        self.assertIn('简历', result['evidence']['品牌识别'])
        self.assertIsNone(app.match('Figma', '')['score'])
        self.assertIsNone(app.match('认真工作', 'Figma')['score'])

    def test_no_short_term_false_positives(self):
        self.assertNotIn('Java', app.keywords('JavaScript'))
        self.assertNotIn('Go', app.keywords('Google'))

    def test_export(self):
        job = {'title': '=HYPERLINK("bad")', 'status': '已投递', 'applied_at': '2026-10-09', 'score': 67, 'keywords': ['Figma']}
        wb = load_workbook(io.BytesIO(app.export([job], '简历：Figma')))
        sheet = wb['秋招进度']
        self.assertEqual(sheet['A2'].data_type, 's')
        self.assertEqual(sheet['E2'].value, '2026-10-09')
        self.assertEqual(sheet['F2'].value, '67%')
        self.assertEqual(sheet['I2'].value, '已投递')
        self.assertEqual(len(sheet.data_validations.dataValidation), 1)
        self.assertIn('个人画像与证据', wb.sheetnames)

    def test_private_url_rejected(self):
        for url in ['http://127.0.0.1', 'http://[::1]', 'file:///etc/passwd', 'http://user:pass@example.com']:
            with self.assertRaises(ValueError):
                app.validate_url(url)

    def test_http_workflow_and_persistence(self):
        with tempfile.TemporaryDirectory() as directory, patch.object(app, 'DB', Path(directory)/'tracker.db'):
            server = app.ThreadingHTTPServer(('127.0.0.1', 0), app.Handler)
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            base = f'http://127.0.0.1:{server.server_port}'
            try:
                self.assertEqual(requests.get(base).status_code, 200)
                text = '品牌设计师\n任职要求：熟悉 Figma、品牌设计、视觉系统，熟悉 Blender 加分。'
                job = requests.post(base+'/api/parse', json={'text': text}).json()
                self.assertEqual(job['title'], '品牌设计师')
                requests.post(base+'/api/profile', json={'profile': '简历：Figma、品牌设计、视觉系统'}).raise_for_status()
                job.update(status='已投递', applied_at='2026-10-09')
                requests.post(base+'/api/save', json=job).raise_for_status()
                state = requests.get(base+'/api/state').json()
                self.assertEqual(len(state['jobs']), 1)
                self.assertEqual(state['jobs'][0]['score'], 75)
                wb = load_workbook(io.BytesIO(requests.get(base+'/api/export').content))
                self.assertEqual(wb.active['A2'].value, '品牌设计师')
                # Fresh database connection confirms committed persisted data.
                with app.connection() as conn:
                    self.assertEqual(len(app.all_jobs(conn, state['profile'])), 1)
                job = state['jobs'][0]
                job['status'] = '一面'
                requests.post(base+'/api/save', json=job).raise_for_status()
                self.assertEqual(requests.get(base+'/api/state').json()['jobs'][0]['status'], '一面')
                requests.post(base+'/api/delete', json={'id': job['id']}).raise_for_status()
                self.assertEqual(requests.get(base+'/api/state').json()['jobs'], [])
            finally:
                server.shutdown()
                server.server_close()


if __name__ == '__main__':
    unittest.main()

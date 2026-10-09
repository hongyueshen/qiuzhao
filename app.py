"""Local autumn recruitment tracker. Run: python app.py."""
import io
import ipaddress
import json
import os
import re
import socket
import sqlite3
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlsplit

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).parent
DB = ROOT / 'data' / 'tracker.db'
STATUSES = ['待投递', '已投递', '笔试', '一面', '二面', '终面', 'Offer', '已拒绝', '已撤回']
TERMS = ['Python', 'Java', 'JavaScript', 'TypeScript', 'C++', 'C#', 'Go', 'SQL', 'MySQL', 'Redis', 'Linux', 'Docker', 'Kubernetes', 'React', 'Vue', 'Spring', 'Git', '算法', '数据结构', '机器学习', '深度学习', '大模型', '数据分析', '数据挖掘', '产品设计', '用户研究', '需求分析', '项目管理', '运营', '市场营销', '内容策划', '沟通', '团队协作', '英语', 'Excel', 'Power BI', 'Tableau', '财务', '会计', '金融', '统计', '本科', '硕士', '博士']
DESIGN_TERMS = {
    '品牌设计': ['品牌设计', 'branding', 'brand design'],
    '视觉设计': ['视觉设计', '视觉传达', 'visual design', 'graphic design'],
    '品牌识别': ['品牌识别', '品牌视觉识别', '品牌形象', 'brand identity', 'VI设计', 'VI 设计'],
    '视觉系统': ['视觉系统', '设计系统', 'design system'],
    '品牌策略': ['品牌策略', '品牌战略', 'brand strategy'],
    '品牌传播': ['品牌传播', '品牌沟通', 'brand communication'],
    '编辑设计': ['编辑设计', '编辑出版', '美术编辑', 'editorial design', '排版'],
    '包装设计': ['包装设计', 'packaging'],
    '字体设计': ['字体设计', 'typography'],
    '动态视觉': ['动态视觉', '动效', '动画', 'motion design', 'motion graphics'],
    '视频剪辑': ['视频剪辑', '视频编辑', '视频后期', 'video editing'],
    '视觉陈列': ['视觉陈列', 'visual merchandising'],
    '社交媒体': ['社交媒体', '新媒体', 'social media'],
    '用户体验': ['用户体验', 'UX', 'user experience'],
    '界面设计': ['界面设计', 'UI', 'user interface'],
    '交互设计': ['交互设计', 'interaction design'],
    '服务设计': ['服务设计', 'service design'],
    '工业设计': ['工业设计', 'industrial design'],
    '产品策划': ['产品策划', '产品规划', 'product planning'],
    '竞品分析': ['竞品分析', 'competitive analysis'],
    '原型设计': ['原型设计', 'prototyping', 'prototype'],
    '作品集': ['作品集', 'portfolio'],
    '创意策划': ['创意策划', '品牌概念策划', 'creative concept'],
    '供应商管理': ['供应商管理', '供应商', 'vendor'],
    '成本控制': ['成本控制', 'cost control'],
    'Photoshop': ['Photoshop', 'PS'],
    'Illustrator': ['Illustrator', 'AI软件'],
    'InDesign': ['InDesign', 'ID软件'],
    'After Effects': ['After Effects', 'AE'],
    'Premiere Pro': ['Premiere Pro', 'Premiere', 'PR软件'],
    'Figma': ['Figma'], 'Sketch': ['Sketch'], 'Blender': ['Blender'],
    'Cinema 4D': ['Cinema 4D', 'C4D'],
    'AI辅助设计': ['AI辅助设计', 'AI 辅助设计', 'AI设计', 'AI 设计', '生成式AI', 'Midjourney', 'ChatGPT'],
    '日语': ['日语', 'Japanese'], '法语': ['法语', 'French'],
}


def contains(text, word):
    pattern = re.escape(word)
    if word.isascii():
        pattern = r'(?<![\w+#])' + pattern + r'(?![\w+#])'
    return bool(re.search(pattern, text, re.I))


def connection():
    DB.parent.mkdir(exist_ok=True)
    conn = sqlite3.connect(DB)
    conn.execute('CREATE TABLE IF NOT EXISTS jobs (id INTEGER PRIMARY KEY, payload TEXT NOT NULL)')
    conn.execute('CREATE TABLE IF NOT EXISTS settings (id INTEGER PRIMARY KEY, profile TEXT NOT NULL)')
    return conn


def keywords(text):
    found = []
    for word in TERMS:
        if contains(text, word):
            found.append(word)
    for term, aliases in DESIGN_TERMS.items():
        if any(contains(text, alias) for alias in aliases) and term not in found:
            found.append(term)
    return found


def match(requirements, profile):
    # Explainable keyword coverage, not a hiring probability.
    required = keywords(requirements)
    available = {x.lower() for x in keywords(profile)}
    hit = [x for x in required if x.lower() in available]
    missing = [x for x in required if x.lower() not in available]
    evidence = {}
    for term in hit:
        aliases = DESIGN_TERMS.get(term, [term])
        for line in profile.splitlines():
            if any(contains(line, alias) for alias in aliases):
                evidence[term] = line[:220]
                break
    return {'score': round(100 * len(hit) / len(required)) if required and profile.strip() else None,
            'matched': hit, 'missing': missing, 'evidence': evidence}


def validate_url(url):
    parsed = urlsplit(url)
    if parsed.scheme not in ('http', 'https') or not parsed.hostname or parsed.username or parsed.password:
        raise ValueError('请输入有效的 http / https 岗位链接')
    if parsed.port and parsed.port not in (80, 443):
        raise ValueError('仅支持标准网页端口')
    host = parsed.hostname.rstrip('.').lower()
    try:
        address = ipaddress.ip_address(host)
    except ValueError:
        address = None
    if address is not None and not address.is_global:
        raise ValueError('仅支持公网招聘网页')
    if address is None and ('.' not in host or host.endswith(('.local', '.localhost', '.internal'))):
        raise ValueError('仅支持公网招聘网页')
    # The Codex egress proxy resolves destinations and enforces private-address
    # blocking upstream. Container DNS is unavailable there; use that supported route.
    if os.environ.get('CODEX_EXEC_SERVER_PROXY_PRIVATE_IPS_VIA_UPSTREAM') and os.environ.get('HTTPS_PROXY'):
        return url
    addresses = socket.getaddrinfo(parsed.hostname, parsed.port or (443 if parsed.scheme == 'https' else 80), type=socket.SOCK_STREAM)
    if not addresses or any(not ipaddress.ip_address(item[4][0]).is_global for item in addresses):
        raise ValueError('仅支持公网招聘网页')
    return url


def fetch(url):
    # Use the environment's supported network proxy; do not load browser cookies.
    with requests.Session() as session:
        for _ in range(6):
            validate_url(url)
            with session.get(url, timeout=(5, 12), allow_redirects=False, stream=True,
                             headers={'User-Agent': 'QiuzhaoTracker/1.0'}) as response:
                if response.is_redirect:
                    from urllib.parse import urljoin
                    url = urljoin(url, response.headers['Location'])
                    continue
                response.raise_for_status()
                if not any(x in response.headers.get('Content-Type', '').lower() for x in ('html', 'text')):
                    raise ValueError('链接未返回可解析的网页，请粘贴岗位正文')
                content = bytearray()
                for chunk in response.iter_content(16384):
                    content.extend(chunk)
                    if len(content) > 2_000_000:
                        raise ValueError('页面过大，请直接粘贴岗位正文')
                return bytes(content)
    raise ValueError('页面重定向次数过多')


def parse_job(html=None, text='', url=''):
    title, company, structured_description = '', '', ''
    if html:
        soup = BeautifulSoup(html, 'html.parser')
        def walk(value):
            if isinstance(value, list):
                for item in value:
                    yield from walk(item)
            elif isinstance(value, dict):
                if value.get('@type') == 'JobPosting' or 'JobPosting' in (value.get('@type') or []):
                    yield value
                for item in value.values():
                    if isinstance(item, (dict, list)):
                        yield from walk(item)
        for script in soup.select('script[type="application/ld+json"]'):
            try:
                posting = next(walk(json.loads(script.get_text())), None)
                if posting:
                    title = str(posting.get('title', ''))
                    organization = posting.get('hiringOrganization') or {}
                    company = str(organization.get('name', '')) if isinstance(organization, dict) else str(organization)
                    structured_description = BeautifulSoup(str(posting.get('description', '')), 'html.parser').get_text('\n', strip=True)
                    break
            except (ValueError, TypeError):
                continue
        if not title:
            heading = soup.find('h1') or soup.find('title')
            title = heading.get_text(' ', strip=True) if heading else ''
        for node in soup(['script', 'style', 'nav', 'footer', 'header', 'noscript']):
            node.decompose()
        text = structured_description or soup.get_text('\n', strip=True)
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    title = title or (lines[0] if lines else '')
    if len(text.strip()) < 20:
        raise ValueError('未获取到完整岗位正文，请复制岗位要求并粘贴解析')
    requirement_start = next((i for i, line in enumerate(lines) if re.search(r'任职要求|岗位要求|职位要求|任职资格|Qualifications|Requirements', line, re.I)), None)
    requirements = '\n'.join(lines[requirement_start:]) if requirement_start is not None else text
    if requirement_start is not None:
        requirements = re.split(r'福利待遇|薪酬福利|申请方式|公司介绍|About us|Benefits', requirements, maxsplit=1, flags=re.I)[0]
    terms = keywords(requirements)
    return {'title': title[:160], 'company': company[:120], 'url': url,
            'requirements': requirements[:12000], 'keywords': terms,
            'warning': '请核对提取结果；网页内容与岗位关键词可能不完整。' if html else '已按粘贴正文解析，请核对岗位名称。'}


def export(jobs, profile=''):
    wb = Workbook()
    sheet = wb.active
    sheet.title = '秋招进度'
    sheet.append(['岗位名称', '公司', '岗位链接', '岗位要求关键词', '投递时间', '关键词匹配度', '已匹配关键词', '材料暂未体现', '投递进度', '备注', '岗位要求原文', '匹配证据'])
    for job in jobs:
        values = [job.get('title', ''), job.get('company', ''), job.get('url', ''), '、'.join(job.get('keywords', [])), job.get('applied_at', ''),
                  f"{job['score']}%" if job.get('score') is not None else '待分析', '、'.join(job.get('matched', [])), '、'.join(job.get('missing', [])), job.get('status', '待投递'), job.get('notes', ''), job.get('requirements', ''), '\n'.join(f'{k}：{v}' for k, v in job.get('evidence', {}).items())]
        # Force text cells so spreadsheet formulas from scraped pages never execute.
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='176B5B')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    widths = [28, 22, 40, 42, 18, 18, 32, 32, 16, 36, 65, 65]
    for col, width in zip('ABCDEFGHIJKL', widths):
        sheet.column_dimensions[col].width = width
    dropdown = DataValidation(type='list', formula1='"' + ','.join(STATUSES) + '"')
    sheet.add_data_validation(dropdown)
    dropdown.add('I2:I10000')
    help_sheet = wb.create_sheet('使用说明')
    for line in ['秋招记录表', '在网页中粘贴岗位链接或正文，核对提取结果后保存。', '匹配度为个人背景与岗位要求的关键词覆盖比例，不代表录取概率。', '未填写个人背景或没有识别到关键词时显示待分析。', '投递进度列支持下拉选择；投递时间请按实际日期填写。', 'Excel 是导出快照，在 Excel 中的修改不会自动同步到网页。']:
        help_sheet.append([line])
    help_sheet.column_dimensions['A'].width = 100
    if profile:
        profile_sheet = wb.create_sheet('个人画像与证据')
        profile_sheet.append(['根据用户提供的简历和作品集整理，请核对并按实际情况更新'])
        for line in profile.splitlines():
            profile_sheet.append([line])
            profile_sheet.cell(profile_sheet.max_row, 1).data_type = 's'
        profile_sheet.column_dimensions['A'].width = 110
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def all_jobs(conn, profile):
    jobs = []
    for key, payload in conn.execute('SELECT id, payload FROM jobs ORDER BY id DESC'):
        job = json.loads(payload)
        job.update(id=key, **match(job.get('requirements', ''), profile))
        jobs.append(job)
    return jobs


class Handler(BaseHTTPRequestHandler):
    def send(self, data, status=200, kind='application/json; charset=utf-8'):
        body = data if isinstance(data, bytes) else json.dumps(data, ensure_ascii=False).encode()
        self.send_response(status)
        self.send_header('Content-Type', kind)
        self.send_header('Content-Length', str(len(body)))
        self.send_header('X-Content-Type-Options', 'nosniff')
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        path = urlsplit(self.path).path
        if path in ('/', '/style.css', '/app.js', '/engine.js', '/terms.json', '/vendor/exceljs.min.js'):
            file = ROOT / 'static' / ('index.html' if path == '/' else path[1:])
            kind = 'application/json; charset=utf-8' if path.endswith('.json') else 'application/javascript; charset=utf-8' if path.endswith('.js') else 'text/css; charset=utf-8' if path.endswith('.css') else 'text/html; charset=utf-8'
            return self.send(file.read_bytes(), kind=kind)
        with connection() as conn:
            row = conn.execute('SELECT profile FROM settings WHERE id=1').fetchone()
            profile = row[0] if row else ''
            jobs = all_jobs(conn, profile)
        if path == '/api/state':
            return self.send({'jobs': jobs, 'profile': profile, 'statuses': STATUSES})
        if path == '/api/export':
            return self.send(export(jobs, profile), kind='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.send({'error': '页面不存在'}, 404)

    def do_POST(self):
        # Reject cross-origin form submissions and bound payload size.
        if 'application/json' not in self.headers.get('Content-Type', ''):
            return self.send({'error': '需要 JSON 请求'}, 415)
        if self.headers.get('Origin') and urlsplit(self.headers['Origin']).netloc != self.headers.get('Host'):
            return self.send({'error': '不允许跨站请求'}, 403)
        try:
            size = int(self.headers.get('Content-Length', '0'))
            if not 0 < size <= 200_000:
                raise ValueError('请求内容为空或过大')
            data = json.loads(self.rfile.read(size))
            if self.path == '/api/parse':
                url = str(data.get('url', '')).strip()
                text = str(data.get('text', '')).strip()
                job = parse_job(text=text, url=url) if text else parse_job(html=fetch(url), url=url)
                return self.send(job)
            with connection() as conn:
                if self.path == '/api/profile':
                    conn.execute('INSERT OR REPLACE INTO settings VALUES (1, ?)', (str(data.get('profile', ''))[:12000],))
                elif self.path == '/api/save':
                    job = {key: str(data.get(key, '')).strip() for key in ('title', 'company', 'url', 'requirements', 'applied_at', 'status', 'notes')}
                    if not job['title']:
                        raise ValueError('请填写岗位名称')
                    if job['status'] not in STATUSES:
                        raise ValueError('投递进度无效')
                    if job['applied_at']:
                        datetime.strptime(job['applied_at'], '%Y-%m-%d')
                    job['keywords'] = keywords(job['requirements'])
                    if data.get('id'):
                        conn.execute('UPDATE jobs SET payload=? WHERE id=?', (json.dumps(job, ensure_ascii=False), int(data['id'])))
                    else:
                        conn.execute('INSERT INTO jobs (payload) VALUES (?)', (json.dumps(job, ensure_ascii=False),))
                elif self.path == '/api/delete':
                    conn.execute('DELETE FROM jobs WHERE id=?', (int(data['id']),))
                else:
                    return self.send({'error': '接口不存在'}, 404)
            self.send({'ok': True})
        except (ValueError, KeyError, TypeError) as exc:
            self.send({'error': str(exc)}, 400)
        except (requests.RequestException, socket.gaierror):
            self.send({'error': '无法读取此招聘网页：可能需要登录、浏览器渲染或限制抓取。请粘贴岗位正文解析。'}, 422)
        except Exception:
            self.send({'error': '处理失败，请重试并检查服务日志'}, 500)


if __name__ == '__main__':
    port = int(os.environ.get('PORT', '8000'))
    server = ThreadingHTTPServer(('127.0.0.1', port), Handler)
    print(f'秋招管理服务启动，端口 {port}', flush=True)
    server.serve_forever()
